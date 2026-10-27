"""Write a small demo corpus as REAL files in 7 formats (md, txt, html, json, csv, pdf, docx)
for manual upload / tests. Derived from the deterministic synthetic generator."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from eval.synthetic import generate


def write_pdf(path: Path, title: str, text: str) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle(title)
    y = 800
    for line in text.splitlines():
        for chunk in [line[i:i + 95] for i in range(0, max(len(line), 1), 95)]:
            if y < 60:
                c.showPage(); y = 800
            c.drawString(40, y, chunk); y -= 14
    c.save()


def main(out: str = "data/sample_files", n_inc: int = 6) -> None:
    out_dir = Path(out); out_dir.mkdir(parents=True, exist_ok=True)
    corpus = generate(n_inc, 0, seed=3)
    ids = {f"{i:04d}" for i in range(1, n_inc + 1)}
    written = 0
    for d in corpus.docs:
        kind, base = d["metadata"].get("kind"), d["doc_id"]
        if kind in ("runbook", "digest", "noise") or base.endswith("-copy") or base.endswith("-v2"):
            continue
        if kind in ("incident_report", "postmortem", "alerts") and base.split("-")[-1] not in ids:
            continue
        fmt, text = d["format"], d["text"]
        p = out_dir / f"{base}.{fmt}"
        if kind == "incident_report":  # show off PDF + DOCX handling
            if int(base.split("-")[-1]) % 2 == 0:
                write_pdf(out_dir / f"{base}.pdf", d["title"], text); written += 1; continue
            doc = docx.Document(); doc.core_properties.title = d["title"]
            for line in text.splitlines():
                if line.startswith("# "): doc.add_heading(line[2:], 1)
                elif line.startswith("## "): doc.add_heading(line[3:], 2)
                elif line.strip(): doc.add_paragraph(line.lstrip("- ").strip() if line.startswith("- ") else line)
            doc.save(str(out_dir / f"{base}.docx")); written += 1; continue
        if fmt == "json":
            obj = dict(l.split(": ", 1) for l in text.splitlines() if ": " in l)
            p.write_text(json.dumps(obj, indent=2)); written += 1; continue
        if fmt == "html":
            body = "".join(f"<h2>{l[2:]}</h2>" if l.startswith("# ") else f"<li>{l[2:]}</li>" if l.startswith("- ") else f"<p>{l}</p>" for l in text.splitlines() if l.strip())
            p.write_text(f"<html><head><title>{d['title']}</title><script>var x=1;</script></head><body>{body}</body></html>"); written += 1; continue
        if fmt == "csv":
            rows = [dict(kv.split(": ", 1) for kv in l.split("; ")) for l in text.splitlines()]
            p.write_text(",".join(rows[0]) + "\n" + "\n".join(",".join(r.get(k, "") for k in rows[0]) for r in rows)); written += 1; continue
        p.write_text(text); written += 1
    print(f"wrote {written} sample files to {out_dir}")
    Path("docs/sample_questions.txt").write_text(
        "What was the severity of INC-0001?\nWhich change caused INC-0003 and who authored it?\nHow long did INC-0002 last?\n"
        "What was the total financial loss caused by INC-0001?\n")


if __name__ == "__main__":
    main(*sys.argv[1:2])
