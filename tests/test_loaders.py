import io
import json

import docx
import pytest
from reportlab.pdfgen import canvas

from app.ingestion.loaders import LoaderError, UnsupportedFormat, content_hash, load_bytes, stable_doc_id


def test_markdown_and_txt():
    d = load_bytes(b"# Heading\n\nHello world 2024-05-06", "a.md")
    assert d.title == "Heading" and d.format == "md" and d.metadata["first_date"] == "2024-05-06"
    assert load_bytes(b"plain text", "b.txt").format == "txt"


def test_html_strips_scripts():
    d = load_bytes(b"<html><head><title>T</title><script>evil()</script></head><body><h1>Hi</h1><p>Body text</p></body></html>", "p.html")
    assert "evil" not in d.text and "Body text" in d.text and d.title == "T"


def test_pdf_roundtrip():
    buf = io.BytesIO(); c = canvas.Canvas(buf); c.setTitle("My PDF"); c.drawString(50, 750, "Severity: SEV1 for INC-0001"); c.save()
    d = load_bytes(buf.getvalue(), "r.pdf")
    assert "SEV1" in d.text and d.metadata["pages"] == 1 and d.format == "pdf"


def test_docx_roundtrip():
    buf = io.BytesIO(); doc = docx.Document(); doc.add_heading("Report", 1); doc.add_paragraph("Root cause: CHG-0001"); doc.save(buf)
    d = load_bytes(buf.getvalue(), "r.docx")
    assert "CHG-0001" in d.text and d.text.startswith("# Report")


def test_json_and_csv():
    j = load_bytes(json.dumps({"id": "CHG-1", "meta": {"author": "Ana"}}).encode(), "c.json")
    assert "meta.author: Ana" in j.text
    c = load_bytes(b"service,alert\npay,High\nledger,Low\n", "a.csv")
    assert "service: pay; alert: High" in c.text and c.metadata["rows"] == 2


@pytest.mark.parametrize("name,data", [("x.pdf", b"not a pdf"), ("x.docx", b"not a zip"), ("x.json", b"{bad"), ("x.txt", b"   \n ")])
def test_corrupt_files_raise_loader_error(name, data):
    with pytest.raises(LoaderError):
        load_bytes(data, name)


def test_unsupported_extension():
    with pytest.raises(UnsupportedFormat):
        load_bytes(b"x", "malware.exe")


def test_ids_and_hash_are_stable():
    assert stable_doc_id("Report 1.md") == stable_doc_id("Report 1.md") != stable_doc_id("Report 2.md")
    assert content_hash("Hello   World") == content_hash("hello world")
