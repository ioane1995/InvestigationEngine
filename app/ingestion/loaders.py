"""Format loaders. Supported: txt, md, html/htm, pdf, docx, json, csv (7 formats; 3 required).
Every loader returns a RawDocument with extracted text + metadata (title, dates, pages, author...)."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any

from app.models import RawDocument

SUPPORTED = {".txt", ".md", ".markdown", ".html", ".htm", ".pdf", ".docx", ".json", ".csv"}
_DATE_RE = re.compile(r"\b(20\d{2}|19\d{2})-(0[1-9]|1[0-2])-(0[1-9]|[12]\d|3[01])\b")


class UnsupportedFormat(ValueError):
    pass


class LoaderError(ValueError):
    """Raised when a file is supported by extension but cannot be parsed (corrupt, encrypted...)."""


def stable_doc_id(source_name: str) -> str:
    """Deterministic id from the file name so re-uploading a file = *update* of the same document."""
    slug = re.sub(r"[^a-zA-Z0-9._-]+", "-", Path(source_name).stem).strip("-").lower() or "doc"
    h = hashlib.sha1(source_name.encode("utf-8")).hexdigest()[:8]
    return f"{slug[:48]}-{h}"


def content_hash(text: str) -> str:
    """Hash of whitespace/case-normalised text -> detects exact duplicate documents."""
    norm = re.sub(r"\s+", " ", text).strip().lower()
    return hashlib.sha256(norm.encode("utf-8")).hexdigest()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _first_heading(text: str) -> str:
    for line in text.splitlines():
        s = line.strip()
        if s:
            return s.lstrip("# ").strip()[:200]
    return ""


def _base_meta(filename: str, data: bytes, fmt: str, text: str) -> dict[str, Any]:
    m: dict[str, Any] = {"filename": filename, "format": fmt, "size_bytes": len(data)}
    d = _DATE_RE.search(text)
    if d:
        m["first_date"] = d.group(0)
    return m


def _load_html(data: bytes) -> tuple[str, str, dict]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    for t in soup(["script", "style", "noscript", "template"]):
        t.decompose()
    title = (soup.title.string or "").strip() if soup.title and soup.title.string else ""
    parts: list[str] = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr", "pre"]):
        if el.name == "tr":
            cells = [c.get_text(" ", strip=True) for c in el.find_all(["th", "td"])]
            line = " | ".join(c for c in cells if c)
        else:
            line = el.get_text(" ", strip=True)
        if not line:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            line = "#" * int(el.name[1]) + " " + line
        parts.append(line)
    text = "\n\n".join(parts) if parts else soup.get_text("\n", strip=True)
    return title or _first_heading(text), text, {}


def _load_pdf(data: bytes) -> tuple[str, str, dict]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise LoaderError("PDF is encrypted")
        pages = [(p.extract_text() or "").strip() for p in reader.pages]
    except PdfReadError as e:
        raise LoaderError(f"invalid PDF: {e}") from e
    text = "\n\n".join(p for p in pages if p)
    if not text:
        raise LoaderError("PDF contains no extractable text (scanned image? OCR is out of scope)")
    info = reader.metadata or {}
    meta: dict[str, Any] = {"pages": len(pages)}
    if info.get("/Author"):
        meta["author"] = str(info.get("/Author"))
    title = str(info.get("/Title") or "") or _first_heading(text)
    return title, text, meta


def _load_docx(data: bytes) -> tuple[str, str, dict]:
    import docx  # python-docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # zipfile.BadZipFile, KeyError ...
        raise LoaderError(f"invalid DOCX: {e}") from e
    lines: list[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading"):
            lvl = re.sub(r"\D", "", style) or "1"
            t = "#" * int(lvl) + " " + t
        elif style == "Title":
            t = "# " + t
        lines.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    text = "\n\n".join(lines)
    cp = d.core_properties
    meta: dict[str, Any] = {}
    if cp.author:
        meta["author"] = cp.author
    title = (cp.title or "").strip() or _first_heading(text)
    return title, text, meta


def _flatten_json(obj: Any, prefix: str = "") -> list[str]:
    lines: list[str] = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            key = f"{prefix}{k}"
            if isinstance(v, (dict, list)):
                lines.extend(_flatten_json(v, key + "."))
            else:
                lines.append(f"{key}: {v}")
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            if isinstance(v, (dict, list)):
                lines.extend(_flatten_json(v, f"{prefix}{i}."))
            else:
                lines.append(f"{prefix.rstrip('.')}[{i}]: {v}")
    else:
        lines.append(f"{prefix.rstrip('.')}: {obj}" if prefix else str(obj))
    return lines


def _load_json(data: bytes) -> tuple[str, str, dict]:
    try:
        obj = json.loads(_decode(data))
    except json.JSONDecodeError as e:
        raise LoaderError(f"invalid JSON: {e}") from e
    title = ""
    if isinstance(obj, dict):
        title = str(obj.get("title") or obj.get("name") or obj.get("id") or "")
    return title, "\n".join(_flatten_json(obj)), {}


def _load_csv(data: bytes) -> tuple[str, str, dict]:
    text = _decode(data)
    try:
        rows = list(csv.DictReader(io.StringIO(text)))
    except csv.Error as e:
        raise LoaderError(f"invalid CSV: {e}") from e
    if not rows:
        raise LoaderError("CSV has no data rows")
    lines = ["; ".join(f"{k}: {v}" for k, v in r.items() if v not in (None, "")) for r in rows]
    return "", "\n".join(lines), {"rows": len(rows), "columns": list(rows[0].keys())}


def load_bytes(data: bytes, filename: str, doc_id: str | None = None) -> RawDocument:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise UnsupportedFormat(f"unsupported file type '{ext}'. Supported: {sorted(SUPPORTED)}")
    fmt = ext.lstrip(".")
    if ext in (".txt", ".md", ".markdown"):
        text = _decode(data)
        title, extra = _first_heading(text), {}
        fmt = "md" if ext != ".txt" else "txt"
    elif ext in (".html", ".htm"):
        title, text, extra = _load_html(data)
        fmt = "html"
    elif ext == ".pdf":
        title, text, extra = _load_pdf(data)
    elif ext == ".docx":
        title, text, extra = _load_docx(data)
    elif ext == ".json":
        title, text, extra = _load_json(data)
    else:
        title, text, extra = _load_csv(data)
    text = text.replace("\x00", "").strip()
    if not text:
        raise LoaderError("document is empty after text extraction")
    meta = _base_meta(filename, data, fmt, text)
    meta.update(extra)
    return RawDocument(
        doc_id=doc_id or stable_doc_id(filename),
        title=(title or Path(filename).stem)[:200],
        text=text,
        source=filename,
        format=fmt,
        metadata=meta,
    )
